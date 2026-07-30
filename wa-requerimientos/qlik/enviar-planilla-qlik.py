"""Planilla semanal de usuarios Qlik.

Toma los pedidos de alta/baja de usuarios acumulados en la Data Table
`qlik_usuarios` de n8n, los agrega al Excel (seccion inferior, FONDO ROJO los
nuevos de la semana; los de semanas anteriores quedan sin color), y envia el
archivo por WhatsApp al celular de la mesa Qlik via el bridge local.

Corre como tarea programada los viernes. Uso manual:
    python enviar-planilla-qlik.py           # envio real
    python enviar-planilla-qlik.py --test    # manda al chat propio (QA)
"""
import json
import os
import shutil
import sys
import urllib.parse
import urllib.request
from datetime import date

import openpyxl
from openpyxl.styles import Font, PatternFill

BASE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(BASE, "template-original.xlsx")
MASTER = os.path.join(BASE, "planilla-master.xlsx")
STATE = os.path.join(BASE, "qlik-lastid.txt")

N8N_KEY = "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJzdWIiOiJiNjBkNWU1YS1lYjNmLTQ4ZTItYTZlOC1iYzUyMTUzZTFhYzMiLCJpc3MiOiJuOG4iLCJhdWQiOiJwdWJsaWMtYXBpIiwianRpIjoiZDFhY2E0YWYtYjExMS00ZmU2LTlkZTItZmIyNWE3ODk1ZDRkIiwiaWF0IjoxNzc2OTUwMjc3fQ.K3a87ubNNw-N_ri3Q7mnYgDD2etP2l9OyoMA9Ut3L5Y"
ROWS_URL = "https://n8n.srv1312396.hstgr.cloud/api/v1/data-tables/8TuBEc9Y7hMJFwjl/rows"

DESTINO = "5491132900699@s.whatsapp.net"  # celular Mesa Qlik
PABLO = "5491157935007@s.whatsapp.net"
BRIDGE_TOKEN_FILE = os.path.expanduser(r"~\dev\whatsapp-mcp-vgp\whatsapp-bridge\store\.bridge-token")
OUTBOX_DIR = os.path.expanduser(r"~\.local\share\whatsapp-mcp\outbox")

ROJO = PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid")


def api_get(url):
    req = urllib.request.Request(url, headers={"X-N8N-API-KEY": N8N_KEY})
    with urllib.request.urlopen(req, timeout=30) as r:
        return json.loads(r.read().decode("utf-8"))


def bridge_send(recipient, message, media_path=None):
    with open(BRIDGE_TOKEN_FILE) as f:
        token = f.read().strip()
    body = {"recipient": recipient, "message": message}
    if media_path:
        body["media_path"] = media_path
    req = urllib.request.Request(
        "http://127.0.0.1:8080/api/send",
        data=json.dumps(body).encode("utf-8"),
        headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"},
        method="POST",
    )
    with urllib.request.urlopen(req, timeout=60) as r:
        return json.loads(r.read().decode("utf-8"))


def main():
    test = "--test" in sys.argv
    destino = PABLO if test else DESTINO

    last_id = 0
    if os.path.exists(STATE):
        last_id = int(open(STATE).read().strip() or 0)

    filtro = urllib.parse.quote(json.dumps({"type": "and", "filters": [{"columnName": "id", "condition": "gt", "value": last_id}]}))
    data = api_get(f"{ROWS_URL}?filter={filtro}").get("data", [])
    nuevos = sorted(data, key=lambda r: r["id"])

    if not os.path.exists(MASTER):
        shutil.copy(TEMPLATE, MASTER)

    wb = openpyxl.load_workbook(MASTER)
    ws = wb.active

    # Los rojos de la semana pasada dejan de ser "nuevos": limpiar fondo
    for row in ws.iter_rows(min_row=15):
        for c in row:
            if c.fill and c.fill.start_color and str(c.fill.start_color.rgb) == "FFFF0000":
                c.fill = PatternFill(fill_type=None)

    # Encabezado de la seccion (una sola vez)
    header_row = None
    for row in ws.iter_rows(min_row=15, max_col=2):
        for c in row:
            if c.value == "PEDIDOS DE USUARIOS":
                header_row = c.row
    if header_row is None:
        header_row = max(ws.max_row + 2, 16)
        ws.cell(row=header_row, column=2, value="PEDIDOS DE USUARIOS").font = Font(bold=True)
        for i, h in enumerate(["FECHA", "PEDIDO", "SOLICITANTE", "REF"]):
            ws.cell(row=header_row + 1, column=2 + i, value=h).font = Font(bold=True)

    # Ultima fila usada de la seccion
    fila = header_row + 2
    while ws.cell(row=fila, column=2).value is not None:
        fila += 1

    for r in nuevos:
        vals = [r.get("fecha", ""), r.get("detalle", ""), r.get("solicitante", ""), r.get("card", "")]
        for i, v in enumerate(vals):
            c = ws.cell(row=fila, column=2 + i, value=v)
            c.fill = ROJO
        fila += 1

    wb.save(MASTER)

    hoy = date.today().isoformat()
    os.makedirs(OUTBOX_DIR, exist_ok=True)
    salida = os.path.join(OUTBOX_DIR, f"Tiempo de Respuesta - Soporte y Mejora {hoy}.xlsx")
    shutil.copy(MASTER, salida)

    msg = (
        f"\U0001F4CA Planilla semanal de usuarios Qlik ({hoy})\n"
        f"Pedidos nuevos esta semana: {len(nuevos)} (marcados en rojo)"
    )
    res = bridge_send(destino, msg, salida)
    print("send:", res)

    if nuevos and res.get("success"):
        with open(STATE, "w") as f:
            f.write(str(nuevos[-1]["id"]))

    aviso = f"✅ Planilla Qlik enviada a {destino} con {len(nuevos)} pedidos nuevos." if res.get("success") else f"⚠️ FALLO el envio de la planilla Qlik: {res}"
    if not test:
        bridge_send(PABLO, aviso)
    print(aviso)


if __name__ == "__main__":
    main()
