import openpyxl, sys

def explore(path, max_rows=20):
    print("=" * 80)
    print("FILE:", path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    print("SHEETS:", wb.sheetnames)
    for name in wb.sheetnames:
        ws = wb[name]
        print("\n" + "-" * 70)
        print(f"SHEET '{name}'  dims={ws.dimensions}  max_row={ws.max_row}  max_col={ws.max_column}")
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            rows.append(row)
            if i >= max_rows:
                break
        # print first rows
        for i, row in enumerate(rows):
            cells = []
            for c in row:
                s = "" if c is None else str(c)
                if len(s) > 22:
                    s = s[:19] + "..."
                cells.append(s)
            print(f"  r{i}: " + " | ".join(cells))
    wb.close()

for p in sys.argv[1:]:
    explore(p)
